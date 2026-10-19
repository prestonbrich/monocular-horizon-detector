import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

HEADER_FONT = Font(bold=True)
HEADER_FILL = PatternFill('solid', fgColor='DDEBF7')
FAIL_FILL = PatternFill('solid', fgColor='FCE4E4')

FRAME_COLUMNS = [
    ('Image', None),
    ('Confidence', '0.00'),
    ('Status', None),
    ('Roll (deg)', '0.00'),
    ('Pitch (deg)', '0.00'),
    ('Offset (px)', '0.0'),
    ('Velocity vx', '0.0000'),
    ('Velocity vy', '0.0000'),
    ('Velocity vz', '0.0000'),
    ('GT Roll (deg)', '0.00'),
    ('GT Pitch (deg)', '0.00'),
    ('GT Offset (px)', '0.0'),
    ('Roll Error (deg)', '0.00'),
    ('Pitch Error (deg)', '0.00'),
    ('Offset Error (px)', '0.0'),
]

ERROR_KEYS = [
    ('Roll error (deg)', 'roll_error_deg'),
    ('Pitch error (deg)', 'pitch_error_deg'),
    ('Offset error (px)', 'offset_error_px'),
]


def _write_header(ws, headers):
    ws.append(headers)
    for cell in ws[1]:
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
    ws.freeze_panes = 'A2'


def _autosize(ws, min_width=10):
    for col in ws.columns:
        width = max(len(str(c.value)) if c.value is not None else 0 for c in col)
        ws.column_dimensions[get_column_letter(col[0].column)].width = max(min_width, width + 2)


def _error_stats(rows, key):
    errors = np.array([abs(r['accuracy'][key]) for r in rows if r['accuracy'] is not None])
    if len(errors) == 0:
        return [None] * 4
    return [float(np.mean(errors)), float(np.median(errors)), float(np.percentile(errors, 90)), len(errors)]


def write_results(rows, failed_frames, output_path, min_confidence=0.75):
    """Write per-frame readings, the confidence-failure tally and accuracy stats to an xlsx file."""
    wb = Workbook()

    # Sheet 1 - every frame: readings, velocity offset (or failure), and accuracy vs ground truth
    ws = wb.active
    ws.title = 'Frames'
    _write_header(ws, [c[0] for c in FRAME_COLUMNS])
    for r in rows:
        v = r['velocity']
        a = r['accuracy'] or {}
        status = 'PASS' if v is not None else 'FAILED - confidence'
        ws.append([
            r['image'], r['confidence'], status, r['roll_deg'], r['pitch_deg'], r['offset_px'],
            *(v if v is not None else (None, None, None)),
            a.get('gt_roll_deg'), a.get('gt_pitch_deg'), a.get('gt_offset_px'),
            a.get('roll_error_deg'), a.get('pitch_error_deg'), a.get('offset_error_px'),
        ])
        if v is None:
            for cell in ws[ws.max_row]:
                cell.fill = FAIL_FILL
    for col_idx, (_, fmt) in enumerate(FRAME_COLUMNS, 1):
        if fmt:
            for row in ws.iter_rows(min_row=2, min_col=col_idx, max_col=col_idx):
                row[0].number_format = fmt
    _autosize(ws)

    # Sheet 2 - summary: confidence tally and accuracy of the pipeline vs ground truth
    ws = wb.create_sheet('Summary')
    total = len(rows)
    passed = [r for r in rows if r['velocity'] is not None]
    no_gt = sum(1 for r in rows if r['accuracy'] is None)

    ws.append(['Confidence Test'])
    ws['A1'].font = HEADER_FONT
    ws.append(['Total frames', total])
    ws.append(['Minimum confidence', min_confidence])
    ws.append(['Frames passed', len(passed)])
    ws.append(['Frames failed (confidence)', len(failed_frames)])
    ws.append(['Failure rate', len(failed_frames) / total if total else None])
    ws.cell(ws.max_row, 2).number_format = '0.0%'
    ws.append(['Frames with no usable ground truth', no_gt])
    ws.append([])

    ws.append(['Accuracy vs Ground Truth (absolute error)', 'Mean', 'Median', '90th percentile', 'Frames'])
    for cell in ws[ws.max_row]:
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
    for group_name, group in (('All frames', rows), ('Passed frames only', passed)):
        ws.append([group_name])
        ws.cell(ws.max_row, 1).font = Font(italic=True)
        for label, key in ERROR_KEYS:
            ws.append(['  ' + label, *_error_stats(group, key)])
            for c in range(2, 5):
                ws.cell(ws.max_row, c).number_format = '0.00'
    _autosize(ws)

    # Sheet 3 - list of frames that failed the confidence test
    ws = wb.create_sheet('Failed Frames')
    _write_header(ws, ['Image', 'Confidence'])
    confidence_by_image = {r['image']: r['confidence'] for r in rows}
    for name in failed_frames:
        ws.append([name, confidence_by_image[name]])
        ws.cell(ws.max_row, 2).number_format = '0.00'
    _autosize(ws)

    wb.save(output_path)
